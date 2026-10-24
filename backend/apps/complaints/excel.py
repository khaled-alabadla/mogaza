"""
Formatted, right-to-left Excel exports (temporary complaints / inquiries, urgent complaints).

openpyxl is imported inside build_workbook: it is only needed when someone exports, and loading it on
every serverless cold start would slow down every other request.
"""
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone

from apps.locations.serializers import visible_user_label

ORGANIZATION = "بلدية غزة — قسم المعلومات والشكاوي"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FONT_NAME = "Arial"
BRAND_GREEN = "1F6F43"
ZEBRA = "EEF6F1"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
DAY_FORMAT = "yyyy-mm-dd"
HEADER_ROW = 3  # 1: title, 2: filters / date / count, 3: column headers


def _local_naive(value):
    """Excel has no time zones: the local (Gaza) wall-clock time."""
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def cell_value(record, field, cell_type, number, viewer):
    """
    The value of one cell. Cell types: "serial" (row number), "text" (kept as text: leading zeros
    survive), "long" (wrapped text), "bool", "choice" (Arabic label), "date" (date + time), "day" (a date),
    "user" (who did it; the hidden developer account is never named), "number".
    """
    if cell_type == "serial":
        return number
    if cell_type == "bool":
        return "نعم" if getattr(record, field) else "لا"
    if cell_type == "choice":
        return getattr(record, f"get_{field}_display")()
    if cell_type == "date":
        return _local_naive(getattr(record, field))
    if cell_type == "day":
        return getattr(record, field)
    if cell_type == "user":
        return visible_user_label(getattr(record, field), viewer) or ""
    value = getattr(record, field)
    if cell_type == "number":
        return value
    return value or ""


def build_workbook(records, *, columns, sheet_name, title, filters, viewer):
    """
    `columns`: [(field, header, width, cell type)]; `filters`: the filter description parts shown under
    the title. Returns the .xlsx bytes.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B7C9BE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    last_col = get_column_letter(len(columns))
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.rightToLeft = True

    # title + subtitle, merged across the table's width
    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    head = ws["A1"]
    head.value = f"{ORGANIZATION}\n{title}"
    head.font = Font(name=FONT_NAME, size=15, bold=True, color=BRAND_GREEN)
    head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 44
    now = timezone.localtime()
    subtitle = ws["A2"]
    subtitle.value = "   |   ".join([*filters, f"تاريخ التصدير: {now:%Y-%m-%d %H:%M}", f"عدد السجلات: {len(records)}"])
    subtitle.data_type = "s"
    subtitle.font = Font(name=FONT_NAME, size=10, color="444444")
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 22

    header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=BRAND_GREEN)
    for index, (_, label, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=index, value=label)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[HEADER_ROW].height = 32

    body_font = Font(name=FONT_NAME, size=10)
    zebra = PatternFill("solid", fgColor=ZEBRA)
    centered = ("serial", "bool", "date", "day", "number", "choice")
    for number, record in enumerate(records, start=1):
        row = HEADER_ROW + number
        for index, (field, _, _, cell_type) in enumerate(columns, start=1):
            value = cell_value(record, field, cell_type, number, viewer)
            cell = ws.cell(row=row, column=index, value=value)
            if isinstance(value, str):
                # always plain text: no "=…" formula (openpyxl's guess) nor number conversion, so
                # leading zeros survive and formula injection is impossible
                cell.data_type = "s"
                cell.number_format = "@"
            elif cell_type == "date":
                cell.number_format = DATETIME_FORMAT
            elif cell_type == "day":
                cell.number_format = DAY_FORMAT
            cell.font, cell.border = body_font, border
            cell.alignment = Alignment(horizontal="center" if cell_type in centered else "right",
                                       vertical="top", wrap_text=True)
            if number % 2 == 0:
                cell.fill = zebra
    last_row = HEADER_ROW + max(len(records), 1)

    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_row}"
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "صفحة &P من &N"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_response(content, arabic_name, ascii_name):
    """The workbook as a download named «<name>_<today>.xlsx» (Arabic name for browsers that support it)."""
    response = HttpResponse(content, content_type=XLSX_TYPE)
    day = timezone.localdate().isoformat()
    response["Content-Disposition"] = (f'attachment; filename="{ascii_name}_{day}.xlsx"; '
                                       f"filename*=UTF-8''{quote(f'{arabic_name}_{day}.xlsx')}")
    return response

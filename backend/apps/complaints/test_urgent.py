import io
from datetime import timedelta

import openpyxl
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.accounts.models import User

from .models import UrgentComplaint

URL = "/api/urgent-complaints/"


class UrgentComplaintTests(APITestCase):
    def setUp(self):
        self.editor = User.objects.create_user("ed", password="pass1234", role=User.Role.EDITOR)
        self.viewer = User.objects.create_user("vw", password="pass1234", role=User.Role.VIEWER)
        self.client.force_authenticate(self.editor)

    def create(self, **data):
        res = self.client.post(URL, data, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        return res.data

    def test_new_complaint_is_unresolved_and_days_become_a_date(self):
        data = self.create(number="١٢٣٤٥", name="أحمد", category="مياه", phone="+970599123456",
                           submitted_days_ago=3, notes="عاجل")
        self.assertEqual(data["status"], "unresolved")
        self.assertEqual(data["status_display"], "لم يتم الحل")
        self.assertEqual(data["number"], "12345")  # Arabic digits normalized
        self.assertEqual(data["phone"], "0599123456")
        self.assertEqual(data["age_days"], 3)
        self.assertEqual(data["submitted_on"], (timezone.localdate() - timedelta(days=3)).isoformat())
        # the age keeps growing: two days later it is 5
        UrgentComplaint.objects.filter(pk=data["id"]).update(submitted_on=timezone.localdate() - timedelta(days=5))
        self.assertEqual(self.client.get(f"{URL}{data['id']}/").data["age_days"], 5)

    def test_empty_record_and_bad_values_are_rejected(self):
        self.assertEqual(self.client.post(URL, {}, format="json").status_code, 400)
        self.assertEqual(self.client.post(URL, {"name": "x", "phone": "123"}, format="json").status_code, 400)
        self.assertEqual(self.client.post(URL, {"name": "x", "submitted_days_ago": -1}, format="json").status_code, 400)

    def test_status_filter_counts_and_mark_resolved(self):
        a = self.create(name="أ", submitted_days_ago=1)
        b = self.create(name="ب", submitted_days_ago=10)
        res = self.client.get(URL, {"status": "unresolved"})
        self.assertEqual([r["name"] for r in res.data["results"]], ["ب", "أ"])  # oldest first
        self.assertEqual(res.data["unresolved_count"], 2)
        self.assertEqual(self.client.post(f"{URL}mark-resolved/", {"ids": [a["id"]]}, format="json").data, {"updated": 1})
        resolved = self.client.get(URL, {"status": "resolved"}).data
        self.assertEqual([r["id"] for r in resolved["results"]], [a["id"]])
        self.assertEqual(resolved["unresolved_count"], 1)
        self.assertEqual(resolved["results"][0]["resolved_by_name"], "ed")
        self.client.post(f"{URL}mark-resolved/", {"ids": [a["id"]], "resolved": False}, format="json")
        self.assertEqual(self.client.get(URL, {"status": "unresolved"}).data["count"], 2)
        self.assertEqual(self.client.get(URL).data["count"], 2)  # «الكل»
        self.assertIsNotNone(b)

    def test_permissions(self):
        self.create(name="أ")
        self.client.force_authenticate(self.viewer)  # viewers don't get this page by default
        self.assertEqual(self.client.get(URL).status_code, 403)
        self.viewer.permissions = ["urgent.view"]
        self.viewer.save()
        self.assertEqual(self.client.get(URL).status_code, 200)
        self.assertEqual(self.client.post(URL, {"name": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.get(f"{URL}export/").status_code, 403)

    def test_excel_export(self):
        self.create(number="00777", name="سعاد", category="صرف صحي", phone="0599000000", submitted_days_ago=4,
                    notes="=HYPERLINK(1)")
        res = self.client.get(f"{URL}export/", {"status": "unresolved"})
        self.assertEqual(res.status_code, 200)
        self.assertIn("urgent_complaints_", res["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertIn("سجل الشكاوى العاجلة", ws["A1"].value)
        self.assertIn("الحالة: لم يتم الحل", ws["A2"].value)
        header = [c.value for c in ws[3]]
        self.assertEqual(header[:8], ["م", "رقم الشكوى", "اسم مقدم الشكوى", "نوع الشكوى", "رقم الجوال", "تاريخ التقديم",
                                      "مقدمة منذ (يوم)", "ملاحظات"])
        row = {label: cell for label, cell in zip(header, ws[4])}
        self.assertEqual(row["رقم الشكوى"].value, "00777")  # leading zeros kept
        self.assertEqual(row["مقدمة منذ (يوم)"].value, 4)
        self.assertEqual(row["ملاحظات"].data_type, "s")  # no formula
        self.assertEqual(len(header), 8)  # no status / resolved date / recorded-by columns

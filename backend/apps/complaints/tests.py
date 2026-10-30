import datetime
from io import BytesIO
from urllib.parse import quote

from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.audit.models import AuditLog

from .models import Complaint

URL = "/api/complaints/"
PASSWORD = "S3cure-Pass-2026"
VALID = {
    "kind": "complaint", "national_id": "401234567", "name": "محمد أحمد", "phone": "0599123456",
    "point": "نقطة الرمال", "building_number": "5A", "street_number": "1050", "category": "انقطاع مياه",
    "address": "تل الهوا - شارع القدس",
}


class ComplaintTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user("admin1", password=PASSWORD, role=User.Role.ADMIN)
        self.editor = User.objects.create_user("editor1", password=PASSWORD, role=User.Role.EDITOR)
        self.viewer = User.objects.create_user("viewer1", password=PASSWORD, role=User.Role.VIEWER)

    def create(self, **changes):
        self.client.force_authenticate(self.editor)
        res = self.client.post(URL, {**VALID, **changes}, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        return res.data

    def test_create_normalizes_and_audits(self):
        data = self.create(national_id=" ٤٠١٢٣٤٥٦٧ ", phone="+970 599-123-456", name="  محمد   أحمد ")
        self.assertEqual((data["national_id"], data["phone"], data["name"]), ("401234567", "0599123456", "محمد أحمد"))
        self.assertEqual((data["kind_display"], data["status"], data["status_display"]),
                         ("شكوى", "pending", "بانتظار الرفع"))
        self.assertEqual(data["building_number"], "5A")
        self.assertEqual(data["created_by_name"], "editor1")
        self.assertTrue(AuditLog.objects.filter(action="CREATE", entity_type="complaint",
                                                entity_id=str(data["id"])).exists())

    def test_inquiry_and_optional_fields(self):
        data = self.create(kind="inquiry", point="", building_number="", street_number="", address="")
        self.assertEqual(data["kind_display"], "استفسار")

    def test_all_data_fields_are_optional(self):
        self.client.force_authenticate(self.editor)
        res = self.client.post(URL, {"phone": "0599123456"}, format="json")  # only a phone number
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual((res.data["kind"], res.data["name"], res.data["national_id"]), ("complaint", "", ""))
        res = self.client.post(URL, {"kind": "inquiry", "category": "استفسار عن موعد المياه"}, format="json")
        self.assertEqual(res.status_code, 201, res.data)

    def test_empty_record_is_rejected(self):
        self.client.force_authenticate(self.editor)
        for payload in ({}, {"kind": "complaint"}, {"name": "   ", "national_id": "", "address": " "}):
            res = self.client.post(URL, payload, format="json")
            self.assertEqual(res.status_code, 400, payload)
            self.assertEqual(str(res.data["non_field_errors"][0]), "أدخل بيانات الشكوى أو الاستفسار: حقل واحد على الأقل.")
        # editing a record cannot empty it either
        pk = self.create()["id"]
        cleared = {field: "" for field in VALID if field != "kind"}
        self.assertEqual(self.client.patch(f"{URL}{pk}/", cleared, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"{URL}{pk}/", {"name": ""}, format="json").status_code, 200)

    def test_filled_fields_are_still_validated(self):
        self.client.force_authenticate(self.editor)
        cases = {
            "national_id": ("12345", "رقم الهوية يجب أن يتكون من 9 أرقام."),
            "phone": ("12345", "رقم الجوال غير صحيح (مثال: 0599123456)."),
            "kind": ("other", "اختر نوع الطلب: شكوى أو استفسار."),
            "building_number": ("5A<script>", "رقم المبنى يحتوي على رموز غير مسموحة."),
        }
        for field, (value, message) in cases.items():
            res = self.client.post(URL, {**VALID, field: value}, format="json")
            self.assertEqual(res.status_code, 400, field)
            self.assertEqual(str(res.data[field][0]), message, field)
        self.assertEqual(Complaint.objects.count(), 0)

    def test_list_filters_search_and_pending_count(self):
        a = self.create(name="سعيد خالد", kind="complaint")
        self.create(name="هالة محمود", kind="inquiry", phone="0569000000")
        self.client.post(f"{URL}mark-uploaded/", {"ids": [a["id"]]}, format="json")
        res = self.client.get(URL, {"upload": "pending"})
        self.assertEqual([r["name"] for r in res.data["results"]], ["هالة محمود"])
        self.assertEqual(res.data["pending_count"], 1)
        self.assertEqual(self.client.get(URL, {"upload": "uploaded"}).data["count"], 1)
        self.assertEqual(self.client.get(URL, {"kind": "inquiry"}).data["count"], 1)
        # the pending badge counts only the selected kind
        self.assertEqual(self.client.get(URL, {"kind": "complaint", "upload": "pending"}).data["pending_count"], 0)
        self.assertEqual(self.client.get(URL, {"kind": "inquiry", "upload": "pending"}).data["pending_count"], 1)
        self.assertEqual(self.client.get(URL, {"search": "0569"}).data["results"][0]["name"], "هالة محمود")
        self.assertEqual(self.client.get(URL, {"search": "سعيد"}).data["count"], 1)

    def test_mark_uploaded_and_back(self):
        ids = [self.create()["id"], self.create(name="ثاني")["id"]]
        res = self.client.post(f"{URL}mark-uploaded/", {"ids": ids}, format="json")
        self.assertEqual(res.data["updated"], 2)
        c = Complaint.objects.get(pk=ids[0])
        self.assertEqual((c.status, c.uploaded_by), ("uploaded", self.editor))
        self.assertIsNotNone(c.uploaded_at)
        # already uploaded: nothing changes
        self.assertEqual(self.client.post(f"{URL}mark-uploaded/", {"ids": ids}, format="json").data["updated"], 0)
        res = self.client.post(f"{URL}mark-uploaded/", {"ids": [ids[0]], "uploaded": False}, format="json")
        self.assertEqual(res.data["updated"], 1)
        c.refresh_from_db()
        self.assertEqual((c.status, c.uploaded_at, c.uploaded_by), ("pending", None, None))
        self.assertEqual(self.client.post(f"{URL}mark-uploaded/", {"ids": []}, format="json").status_code, 400)

    def test_edit_and_soft_delete(self):
        pk = self.create()["id"]
        res = self.client.patch(f"{URL}{pk}/", {"category": "تسرب مياه"}, format="json")
        self.assertEqual(res.data["category"], "تسرب مياه")
        self.assertEqual(self.client.delete(f"{URL}{pk}/").status_code, 204)
        self.assertEqual(self.client.get(URL).data["count"], 0)
        self.assertFalse(Complaint.objects.get(pk=pk).is_active)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.post(f"{URL}{pk}/restore/").status_code, 200)

    def export(self, **params):
        res = self.client.get(f"{URL}export/", params)
        self.assertEqual(res.status_code, 200, getattr(res, "data", None))
        return res, load_workbook(BytesIO(res.content)).active

    def test_export_xlsx_complaints(self):
        self.create(name="=HYPERLINK(1)", national_id="012345678", phone="0599123456", channel="instagram",
                    area="ليست للشكوى", details="سطر 1\nسطر 2")
        self.create(name="هالة", kind="inquiry")
        res, ws = self.export(kind="complaint")
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        disposition = res["Content-Disposition"]
        day = timezone.localdate().isoformat()
        self.assertIn(f'filename="complaints_{day}.xlsx"', disposition)
        self.assertIn("filename*=UTF-8''" + quote(f"شكاوى_{day}.xlsx"), disposition)
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertIn("بلدية غزة", ws["A1"].value)
        self.assertIn("سجل الشكاوى المؤقتة", ws["A1"].value)
        self.assertIn("عدد السجلات: 1", ws["A2"].value)
        header = [c.value for c in ws[3]]
        self.assertEqual(header, ["م", "إخفاء بيانات صاحب الشكوى", "رقم الهوية", "الاسم", "رقم الجوال", "رقم النقطة",
                                  "رقم المبنى", "رقم الشارع", "نوع الشكوى", "المصدر", "العنوان",
                                  "تفاصيل الشكوى", "الحالة", "تاريخ التسجيل", "سجّلها", "تاريخ الرفع"])
        self.assertNotIn("الحي / المنطقة", header)  # not on the complaint form any more
        self.assertNotIn("الاشتراك", header)
        self.assertEqual(ws.freeze_panes, "A4")
        self.assertEqual(ws.max_row, 4)  # one complaint, the inquiry is not exported
        row = {label: cell for label, cell in zip(header, ws[4])}
        self.assertEqual(row["م"].value, 1)
        self.assertEqual(row["رقم الهوية"].value, "012345678")  # leading zero kept, as text
        self.assertEqual(row["رقم الهوية"].data_type, "s")
        self.assertEqual(row["رقم الجوال"].value, "0599123456")
        self.assertEqual(row["الاسم"].value, "=HYPERLINK(1)")
        self.assertEqual(row["الاسم"].data_type, "s")  # plain text, not a formula
        self.assertEqual(row["المصدر"].value, "انستقرام")
        self.assertEqual(row["إخفاء بيانات صاحب الشكوى"].value, "لا")
        self.assertEqual(row["تفاصيل الشكوى"].value, "سطر 1\nسطر 2")
        self.assertEqual(row["سجّلها"].value, "editor1")
        self.assertIsInstance(row["تاريخ التسجيل"].value, datetime.datetime)
        self.assertIsNone(row["تاريخ الرفع"].value)

    def test_export_xlsx_inquiries_and_filters(self):
        self.create(name="سعيد", kind="inquiry", area="الرمال", channel="visit")
        self.create(name="هالة", kind="inquiry")
        self.create(name="شكوى")
        res, ws = self.export(kind="inquiry", search="سعيد", upload="pending")
        self.assertIn("filename*=UTF-8''" + quote("استفسارات_"), res["Content-Disposition"])
        self.assertIn("سجل الاستفسارات", ws["A1"].value)
        self.assertIn("سعيد", ws["A2"].value)
        header = [c.value for c in ws[3]]
        self.assertEqual(header, ["م", "رقم الهوية", "الاسم", "رقم الجوال", "نوع الاستفسار", "المصدر", "الحي",
                                  "نص الاستفسار", "الحالة", "تاريخ التسجيل", "سجّلها", "تاريخ الرفع"])
        self.assertEqual(ws.max_row, 4)
        row = {label: cell.value for label, cell in zip(header, ws[4])}
        self.assertEqual((row["الاسم"], row["الحي"], row["المصدر"], row["الحالة"]),
                         ("سعيد", "الرمال", "زيارة", "بانتظار الرفع"))

    def test_export_requires_kind(self):
        self.client.force_authenticate(self.editor)
        for params in ({}, {"kind": "other"}):
            res = self.client.get(f"{URL}export/", params)
            self.assertEqual(res.status_code, 400)
            self.assertIn("kind", res.data)

    def test_export_hides_the_superuser_name(self):
        root = User.objects.create_superuser("root", "root@example.com", PASSWORD, first_name="المطور")
        self.client.force_authenticate(root)
        pk = self.client.post(URL, VALID, format="json").data["id"]
        self.client.post(f"{URL}mark-uploaded/", {"ids": [pk]}, format="json")
        self.client.force_authenticate(self.admin)
        _, ws = self.export(kind="complaint")
        values = [c.value for c in ws[4]]
        self.assertNotIn("المطور", values)
        self.assertNotIn("root", values)
        data = self.client.get(f"{URL}{pk}/").data
        self.assertEqual((data["created_by_name"], data["updated_by_name"], data["uploaded_by_name"]),
                         (None, None, None))
        # the superuser still sees its own name
        self.client.force_authenticate(root)
        _, ws = self.export(kind="complaint")
        self.assertIn("المطور", [c.value for c in ws[4]])
        self.assertEqual(self.client.get(f"{URL}{pk}/").data["created_by_name"], "المطور")

    @override_settings(PUBLIC_SEARCH_ENABLED=True)
    def test_viewers_and_anonymous_have_no_access(self):
        pk = self.create()["id"]
        for user in (self.viewer, None):
            self.client.force_authenticate(user)
            for res in (self.client.get(URL), self.client.get(f"{URL}{pk}/"), self.client.get(f"{URL}export/"),
                        self.client.post(URL, VALID, format="json"),
                        self.client.post(f"{URL}mark-uploaded/", {"ids": [pk]}, format="json"),
                        self.client.delete(f"{URL}{pk}/")):
                self.assertEqual(res.status_code, 403)


class InquiryAndComplaintFieldsTests(APITestCase):
    """A complaint and an inquiry carry the same fields as the main system's two forms."""

    def setUp(self):
        self.client.force_authenticate(User.objects.create_user("ed", password=PASSWORD, role=User.Role.EDITOR))

    def test_inquiry_fields(self):
        res = self.client.post(URL, {
            "kind": "inquiry", "national_id": "401234567", "name": "سعيد", "phone": "0599123456",
            "category": "استفسار عن فاتورة", "channel": "phone", "area": "الرمال الجنوبي",
            "details": "  السطر الأول   \n\n  السطر   الثاني  \n",
        }, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual((res.data["kind_display"], res.data["channel"], res.data["channel_display"], res.data["area"]),
                         ("استفسار", "phone", "اتصال هاتفي", "الرمال الجنوبي"))
        self.assertEqual(res.data["details"], "السطر الأول\n\nالسطر الثاني")  # line breaks kept, spaces cleaned
        self.assertEqual(self.client.get(URL, {"search": "فاتوره"}).data["count"], 1)  # Arabic-normalized search

    def test_complaint_fields(self):
        res = self.client.post(URL, {
            "kind": "complaint", "hide_identity": True, "national_id": "401234567", "name": "محمد", "point": "12",
            "phone": "0599123456", "building_number": "5A", "street_number": "1050", "subscription": "77-123",
            "category": "تسرب مياه", "channel": "whatsapp", "area": "الدرج", "address": "شارع الوحدة",
            "details": "منذ يومين",
        }, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        for field, value in (("hide_identity", True), ("subscription", "77-123"), ("channel_display", "واتساب"),
                             ("details", "منذ يومين")):
            self.assertEqual(res.data[field], value, field)
        for removed in ("submitter_type", "receipt_prefix", "receipt_number", "exception_reason", "title"):
            self.assertNotIn(removed, res.data)
        # the source is searchable by its Arabic label
        self.assertEqual(self.client.get(URL, {"search": "واتساب"}).data["count"], 1)

    def test_codes_and_source_are_validated(self):
        res = self.client.post(URL, {"kind": "complaint", "subscription": "12<34>", "name": "x"}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(str(res.data["subscription"][0]), "الاشتراك يحتوي على رموز غير مسموحة.")
        res = self.client.post(URL, {"kind": "inquiry", "channel": "email", "name": "x"}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(str(res.data["channel"][0]), "اختر المصدر من القائمة.")
        # the «إخفاء البيانات» flag alone is not a complaint
        res = self.client.post(URL, {"kind": "complaint", "hide_identity": True}, format="json")
        self.assertEqual(res.status_code, 400)
        res = self.client.post(URL, {"kind": "complaint", "address": "ن" * 301}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("address", res.data)

    def test_all_sources(self):
        expected = [("facebook", "فيسبوك"), ("twitter", "تويتر"), ("whatsapp", "واتساب"), ("website", "موقع البلدية"),
                    ("phone", "اتصال هاتفي"), ("app", "تطبيق البلدية"), ("visit", "زيارة"), ("instagram", "انستقرام")]
        self.assertEqual(list(Complaint.Channel.choices), expected)
        for code, label in expected:
            res = self.client.post(URL, {"kind": "inquiry", "channel": code, "name": "x"}, format="json")
            self.assertEqual(res.status_code, 201, res.data)
            self.assertEqual(res.data["channel_display"], label)


class ComplaintPermissionTests(APITestCase):
    """Per-user permissions: complaints.* and inquiries.* are separate."""

    def setUp(self):
        editor = User.objects.create_user("ed", password=PASSWORD, role=User.Role.EDITOR)
        self.client.force_authenticate(editor)
        self.complaint = self.client.post(URL, VALID, format="json").data["id"]
        self.inquiry = self.client.post(URL, {"kind": "inquiry", "name": "سؤال"}, format="json").data["id"]

    def as_user(self, permissions, role=User.Role.EDITOR):
        user = User.objects.create_user(f"u{User.objects.count()}", password=PASSWORD, role=role,
                                        permissions=permissions)
        self.client.force_authenticate(user)
        return user

    def test_only_inquiries_view(self):
        self.as_user(["inquiries.view"])
        res = self.client.get(URL)
        self.assertEqual([r["id"] for r in res.data["results"]], [self.inquiry])
        self.assertEqual(res.data["pending_count"], 1)
        self.assertEqual(self.client.get(URL, {"kind": "complaint"}).data["count"], 0)
        self.assertEqual(self.client.get(f"{URL}{self.complaint}/").status_code, 404)
        self.assertEqual(self.client.get(f"{URL}{self.inquiry}/").status_code, 200)
        # no edit permission at all
        self.assertEqual(self.client.post(URL, {"kind": "inquiry", "name": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.patch(f"{URL}{self.inquiry}/", {"name": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.get(f"{URL}export/", {"kind": "inquiry"}).status_code, 403)

    def test_inquiries_editor_cannot_touch_complaints(self):
        self.as_user(["inquiries.view", "inquiries.edit", "complaints.view"])
        self.assertEqual(self.client.post(URL, VALID, format="json").status_code, 403)  # a complaint
        self.assertEqual(self.client.post(URL, {"kind": "inquiry", "name": "x"}, format="json").status_code, 201)
        self.assertEqual(self.client.patch(f"{URL}{self.complaint}/", {"name": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.delete(f"{URL}{self.complaint}/").status_code, 403)
        # an inquiry cannot be turned into a complaint either
        self.assertEqual(self.client.patch(f"{URL}{self.inquiry}/", {"kind": "complaint"}, format="json").status_code,
                         403)
        res = self.client.post(f"{URL}mark-uploaded/", {"ids": [self.inquiry, self.complaint]}, format="json")
        self.assertEqual(res.status_code, 403)
        self.assertEqual(Complaint.objects.filter(status="uploaded").count(), 0)  # nothing half-done
        res = self.client.post(f"{URL}mark-uploaded/", {"ids": [self.inquiry]}, format="json")
        self.assertEqual(res.data["updated"], 1)
        self.assertEqual(self.client.delete(f"{URL}{self.inquiry}/").status_code, 204)

    def test_export_needs_the_kinds_export_permission(self):
        self.as_user(["complaints.view", "complaints.export", "inquiries.view"])
        self.assertEqual(self.client.get(f"{URL}export/", {"kind": "complaint"}).status_code, 200)
        self.assertEqual(self.client.get(f"{URL}export/", {"kind": "inquiry"}).status_code, 403)

    def test_viewer_defaults_and_empty_permissions(self):
        for permissions, role in ((None, User.Role.VIEWER), ([], User.Role.EDITOR)):
            self.as_user(permissions, role)
            self.assertEqual(self.client.get(URL).status_code, 403)
            self.assertEqual(self.client.get(f"{URL}export/", {"kind": "complaint"}).status_code, 403)


class ComplaintPointTests(APITestCase):
    """«شكوى على الخريطة»: a complaint's number and type where it is."""

    URL = "/api/complaint-points/"
    BODY = {"number": " 2026/155 ", "category": "تسرب مياه", "latitude": "31.512345", "longitude": "34.456789"}

    def setUp(self):
        self.editor = User.objects.create_user("ed", password=PASSWORD, role=User.Role.EDITOR,
                                               permissions=["complaints.view", "complaints.edit"])
        self.reader = User.objects.create_user("rd", password=PASSWORD, role=User.Role.EDITOR,
                                               permissions=["complaints.view"])

    def test_add_list_edit_delete_with_audit(self):
        self.client.force_authenticate(self.editor)
        res = self.client.post(self.URL, self.BODY, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual((res.data["number"], res.data["category"]), ("2026/155", "تسرب مياه"))
        pk = res.data["id"]
        self.assertEqual(self.client.patch(f"{self.URL}{pk}/", {"category": "انقطاع مياه"}, format="json").data["category"],
                         "انقطاع مياه")
        self.client.force_authenticate(self.reader)
        self.assertEqual([p["number"] for p in self.client.get(self.URL).data], ["2026/155"])
        self.assertEqual(self.client.delete(f"{self.URL}{pk}/").status_code, 403)  # reading only
        self.client.force_authenticate(self.editor)
        self.assertEqual(self.client.delete(f"{self.URL}{pk}/").status_code, 204)
        actions = sorted(AuditLog.objects.filter(entity_type="complaintpoint").values_list("action", flat=True))
        self.assertEqual(actions, ["CREATE", "DELETE", "UPDATE"])

    def test_number_type_and_place_required(self):
        self.client.force_authenticate(self.editor)
        for change in ({"number": " "}, {"category": ""}, {"latitude": "31.9"}):
            res = self.client.post(self.URL, {**self.BODY, **change}, format="json")
            self.assertEqual(res.status_code, 400, change)
        res = self.client.post(self.URL, {"number": "1", "category": "x"}, format="json")
        self.assertEqual(res.status_code, 400)

    def test_without_complaints_view_nothing(self):
        other = User.objects.create_user("no", password=PASSWORD, role=User.Role.EDITOR, permissions=["inquiries.view"])
        self.client.force_authenticate(other)
        self.assertEqual(self.client.get(self.URL).status_code, 403)

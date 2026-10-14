"""
Temporary complaints log («الشكاوى المؤقتة»), used while the main complaints system is down:

    GET/POST            /api/complaints/            ?upload=pending|uploaded  ?kind=complaint|inquiry  ?search=
                        (a complaint and an inquiry have different fields, as in the main system's forms)
    GET/PUT/PATCH/DELETE /api/complaints/{id}/       (DELETE = soft delete, admin can restore)
    POST                /api/complaints/mark-uploaded/   {"ids": [...], "uploaded": true|false}
    GET                 /api/complaints/export/      ?kind= (required) + the list's filters → Excel (.xlsx)

Per-user permissions, by kind (the records hold citizens' ID and phone numbers):
complaints.view / .edit / .export and inquiries.view / .edit / .export. A user only ever sees the kinds
they may view; adding, editing, deleting and marking «تم الرفع» need «<kind>.edit» for every record touched.
"""
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.permissions import SAFE_METHODS, BasePermission
from rest_framework.response import Response

from apps.accounts.permissions import DENIED, has_perm_key
from apps.audit.models import AuditLog
from apps.common.params import text_param
from apps.locations.serializers import visible_user_label
from apps.locations.views import TrackedViewSet

from .models import Complaint
from .serializers import ComplaintSerializer, MarkUploadedSerializer

KIND_REQUIRED = "اختر نوع السجل المراد تصديره: شكاوى (complaint) أو استفسارات (inquiry)."
KINDS = Complaint.Kind.values


def allowed_kinds(user, action):
    """The kinds whose «<kind>.<action>» permission the user holds."""
    return [kind for kind in KINDS if has_perm_key(user, Complaint.perm_key(kind, action))]


class ComplaintPermission(BasePermission):
    """Coarse gate (some kind allowed); the view then checks the kind of every record touched."""

    message = DENIED

    def has_permission(self, request, view):
        if getattr(view, "action", None) == "export":
            needed = "export"
        elif request.method in SAFE_METHODS:
            needed = "view"
        else:
            needed = "edit"
        return bool(allowed_kinds(request.user, needed))


# --- Excel export ---------------------------------------------------------------------------------------------
# Columns of each kind, in the order of that kind's form. (field, header, width, cell type)
# Cell types: "serial" (row number), "text" (stored as text: leading zeros kept), "long" (wrapped text),
# "bool", "choice" (Arabic label), "date", "user".
EXPORT_COLUMNS = {
    Complaint.Kind.COMPLAINT: [
        ("serial", "م", 6, "serial"),
        ("hide_identity", "إخفاء بيانات صاحب الشكوى", 14, "bool"),
        ("national_id", "رقم الهوية", 14, "text"),
        ("name", "الاسم", 24, "text"),
        ("phone", "رقم الجوال", 14, "text"),
        ("point", "رقم النقطة", 14, "text"),
        ("building_number", "رقم المبنى", 12, "text"),
        ("street_number", "رقم الشارع", 12, "text"),
        ("category", "نوع الشكوى", 20, "text"),
        ("channel", "المصدر", 14, "choice"),
        ("address", "العنوان", 28, "long"),
        ("details", "تفاصيل الشكوى", 45, "long"),
        ("status", "الحالة", 14, "choice"),
        ("created_at", "تاريخ التسجيل", 18, "date"),
        ("created_by", "سجّلها", 18, "user"),
        ("uploaded_at", "تاريخ الرفع", 18, "date"),
    ],
    Complaint.Kind.INQUIRY: [
        ("serial", "م", 6, "serial"),
        ("national_id", "رقم الهوية", 14, "text"),
        ("name", "الاسم", 24, "text"),
        ("phone", "رقم الجوال", 14, "text"),
        ("category", "نوع الاستفسار", 22, "text"),
        ("channel", "المصدر", 14, "choice"),
        ("area", "الحي", 18, "text"),
        ("details", "نص الاستفسار", 50, "long"),
        ("status", "الحالة", 14, "choice"),
        ("created_at", "تاريخ التسجيل", 18, "date"),
        ("created_by", "سجّلها", 18, "user"),
        ("uploaded_at", "تاريخ الرفع", 18, "date"),
    ],
}
EXPORT_TITLES = {Complaint.Kind.COMPLAINT: "سجل الشكاوى المؤقتة", Complaint.Kind.INQUIRY: "سجل الاستفسارات"}
EXPORT_FILE_NAMES = {Complaint.Kind.COMPLAINT: ("شكاوى", "complaints"),
                     Complaint.Kind.INQUIRY: ("استفسارات", "inquiries")}
ORGANIZATION = "بلدية غزة — قسم المعلومات والشكاوي"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FONT_NAME = "Arial"
BRAND_GREEN = "1F6F43"
ZEBRA = "EEF6F1"
THIN = Side(style="thin", color="B7C9BE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FORMAT = "yyyy-mm-dd hh:mm"
HEADER_ROW = 3  # 1: title, 2: filters / date / count, 3: column headers


def _local_naive(value):
    """Excel has no time zones: the local (Gaza) wall-clock time."""
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def _filter_description(request):
    parts = []
    upload = text_param(request, "upload")
    parts.append(f"الحالة: {Complaint.Status(upload).label}" if upload in Complaint.Status.values
                 else "الحالة: الكل")
    search = text_param(request, "search")
    if search:
        parts.append(f"البحث: «{search[:100]}»")
    return parts


def build_workbook(records, kind, request):
    """The records of one kind as a formatted, right-to-left Excel workbook (bytes)."""
    columns = EXPORT_COLUMNS[kind]
    last_col = get_column_letter(len(columns))
    wb = Workbook()
    ws = wb.active
    ws.title = EXPORT_FILE_NAMES[kind][0]
    ws.sheet_view.rightToLeft = True

    # title + subtitle, merged across the table's width
    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    title = ws["A1"]
    title.value = f"{ORGANIZATION}\n{EXPORT_TITLES[kind]}"
    title.font = Font(name=FONT_NAME, size=15, bold=True, color=BRAND_GREEN)
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 44
    now = timezone.localtime()
    subtitle = ws["A2"]
    subtitle.value = "   |   ".join([*_filter_description(request),
                                     f"تاريخ التصدير: {now:%Y-%m-%d %H:%M}", f"عدد السجلات: {len(records)}"])
    subtitle.data_type = "s"
    subtitle.font = Font(name=FONT_NAME, size=10, color="444444")
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 22

    header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=BRAND_GREEN)
    for index, (_, label, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=index, value=label)
        cell.font, cell.fill, cell.border = header_font, header_fill, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[HEADER_ROW].height = 32

    body_font = Font(name=FONT_NAME, size=10)
    zebra = PatternFill("solid", fgColor=ZEBRA)
    viewer = request.user
    for number, record in enumerate(records, start=1):
        row = HEADER_ROW + number
        for index, (field, _, _, cell_type) in enumerate(columns, start=1):
            if cell_type == "serial":
                value = number
            elif cell_type == "bool":
                value = "نعم" if getattr(record, field) else "لا"
            elif cell_type == "choice":
                value = getattr(record, f"get_{field}_display")()
            elif cell_type == "date":
                value = _local_naive(getattr(record, field))
            elif cell_type == "user":
                value = visible_user_label(getattr(record, field), viewer) or ""
            else:
                value = getattr(record, field) or ""
            cell = ws.cell(row=row, column=index, value=value)
            if isinstance(value, str):
                # always plain text: no "=…" formula (openpyxl's guess) nor number conversion, so
                # leading zeros survive and formula injection is impossible
                cell.data_type = "s"
                cell.number_format = "@"
            elif cell_type == "date":
                cell.number_format = DATE_FORMAT
            cell.font, cell.border = body_font, BORDER
            cell.alignment = Alignment(horizontal="center" if cell_type in ("serial", "bool", "date") else "right",
                                       vertical="top", wrap_text=True)
            if number % 2 == 0:
                cell.fill = zebra
    last_row = HEADER_ROW + max(len(records), 1)

    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_row}"
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "صفحة &P من &N"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ComplaintViewSet(TrackedViewSet):
    model = Complaint
    serializer_class = ComplaintSerializer
    permission_classes = [ComplaintPermission]
    primary_name_field = "name"

    def get_queryset(self):
        # only the kinds this user may see (others answer 404, as if they did not exist)
        qs = (super().get_queryset().select_related("uploaded_by")
              .filter(kind__in=allowed_kinds(self.request.user, "view")))
        upload = text_param(self.request, "upload")
        if upload in Complaint.Status.values:
            qs = qs.filter(status=upload)
        kind = text_param(self.request, "kind")
        if kind in KINDS:
            qs = qs.filter(kind=kind)
        return qs

    def check_edit(self, *kinds):
        for kind in kinds:
            if not has_perm_key(self.request.user, Complaint.perm_key(kind, "edit")):
                raise PermissionDenied(DENIED)

    def perform_create(self, serializer):
        self.check_edit(serializer.validated_data.get("kind", Complaint.Kind.COMPLAINT))
        super().perform_create(serializer)

    def perform_update(self, serializer):
        self.check_edit(serializer.instance.kind, serializer.validated_data.get("kind", serializer.instance.kind))
        super().perform_update(serializer)

    def perform_destroy(self, instance):
        self.check_edit(instance.kind)
        super().perform_destroy(instance)

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        # how many (of the selected kind) are still waiting to be entered into the main system (tab badge)
        pending = Complaint.active.filter(status=Complaint.Status.PENDING,
                                          kind__in=allowed_kinds(request.user, "view"))
        kind = text_param(request, "kind")
        if kind in KINDS:
            pending = pending.filter(kind=kind)
        response.data["pending_count"] = pending.count()
        return response

    @action(detail=False, methods=["post"], url_path="mark-uploaded")
    def mark_uploaded(self, request):
        """Mark complaints as entered into the main system (or back to pending). 403 if any is not editable."""
        data = MarkUploadedSerializer(data=request.data)
        data.is_valid(raise_exception=True)
        uploaded = data.validated_data["uploaded"]
        target = Complaint.Status.UPLOADED if uploaded else Complaint.Status.PENDING
        records = list(Complaint.active.filter(pk__in=data.validated_data["ids"],
                                               kind__in=allowed_kinds(request.user, "view")))
        self.check_edit(*{record.kind for record in records})
        changed = 0
        for complaint in records:
            if complaint.status == target:
                continue
            before = complaint.snapshot()
            complaint.status = target
            complaint.uploaded_at = timezone.now() if uploaded else None
            complaint.uploaded_by = request.user if uploaded else None
            complaint.updated_by = request.user
            complaint.save(update_fields=["status", "uploaded_at", "uploaded_by", "updated_by", "updated_at"])
            self.audit(AuditLog.Action.UPDATE, complaint, before=before, after=complaint.snapshot())
            changed += 1
        return Response({"updated": changed})

    @action(detail=False, methods=["get"])
    def export(self, request):
        """The filtered list of one kind (all pages) as a formatted Excel workbook."""
        kind = text_param(request, "kind")
        if kind not in KINDS:
            return Response({"kind": [KIND_REQUIRED]}, status=status.HTTP_400_BAD_REQUEST)
        if not has_perm_key(request.user, Complaint.perm_key(kind, "export")):
            raise PermissionDenied(DENIED)
        records = list(self.get_queryset().select_related("created_by").order_by("created_at", "id"))
        content = build_workbook(records, kind, request)
        response = HttpResponse(content, content_type=XLSX_TYPE)
        arabic, ascii_name = EXPORT_FILE_NAMES[kind]
        day = timezone.localdate().isoformat()
        response["Content-Disposition"] = (f'attachment; filename="{ascii_name}_{day}.xlsx"; '
                                           f"filename*=UTF-8''{quote(f'{arabic}_{day}.xlsx')}")
        return response

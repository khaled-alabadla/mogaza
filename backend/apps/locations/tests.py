import tempfile
from io import StringIO
from pathlib import Path

from django.core.management import call_command
from django.test import override_settings
from rest_framework import status
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.audit.models import AuditLog

from .models import Location, StreetName
from .text import normalize_for_search

URL = "/api/locations/"


def make_user(username, role, password="S3cure-Pass-2026"):
    return User.objects.create_user(username=username, password=password, role=role)


class LocationTestBase(APITestCase):
    def setUp(self):
        self.admin = make_user("admin1", User.Role.ADMIN)
        self.editor = make_user("editor1", User.Role.EDITOR)
        self.viewer = make_user("viewer1", User.Role.VIEWER)
        self.council = Location.objects.create(place_name="مجلس الوزراء", description="تل الهوا / شارع القدس",
                                               building_number="95", street_number="1050")
        Location.objects.create(place_name="مسجد المحطة", description="مسجد عمر بن عبدالعزيز",
                                building_number="39", street_number="8159")
        Location.objects.create(place_name="مسجد حمزة", description="مسجد سيد شهداء غزة",
                                building_number="189", street_number="8000")
        Location.objects.create(place_name="مدرسة بنات الصبرة", description="مدرسة الصبرة الإعدادية المشتركة",
                                building_number="55F", street_number="2456")
        Location.objects.create(place_name="مسجد احمد ياسين", description="مسجد أحمد ياسين ( مخيم الشاطئ)",
                                building_number="3", street_number="60805")

    def search(self, query, user=None):
        if user:
            self.client.force_authenticate(user)
        return self.client.get(URL, {"search": query})

    def names(self, response):
        return [r["place_name"] for r in response.data["results"]]


class SearchTests(LocationTestBase):
    def test_exact_search(self):
        res = self.search("مجلس الوزراء", self.viewer)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self.names(res), ["مجلس الوزراء"])
        row = res.data["results"][0]
        self.assertEqual((row["description"], row["building_number"], row["street_number"]),
                         ("تل الهوا / شارع القدس", "95", "1050"))

    def test_partial_search_prefix_and_suffix(self):
        self.assertIn("مجلس الوزراء", self.names(self.search("مجلس", self.viewer)))
        self.assertIn("مجلس الوزراء", self.names(self.search("الوزراء", self.viewer)))
        self.assertIn("مجلس الوزراء", self.names(self.search("وزر", self.viewer)))

    def test_multiple_results_returned(self):
        names = self.names(self.search("مسجد", self.viewer))
        self.assertEqual(len(names), 3)
        self.assertEqual(set(names), {"مسجد المحطة", "مسجد حمزة", "مسجد احمد ياسين"})

    def test_search_in_description(self):
        self.assertEqual(self.names(self.search("تل الهوا", self.viewer)), ["مجلس الوزراء"])

    def test_search_by_numbers(self):
        self.assertIn("مدرسة بنات الصبرة", self.names(self.search("55F", self.viewer)))
        self.assertIn("مدرسة بنات الصبرة", self.names(self.search("55f", self.viewer)))

    def test_arabic_normalization(self):
        # hamza variants, taa marbuta, and Arabic-Indic digits
        self.assertIn("مسجد احمد ياسين", self.names(self.search("أحمد", self.viewer)))
        self.assertIn("مدرسة بنات الصبرة", self.names(self.search("مدرسه", self.viewer)))
        self.assertIn("مجلس الوزراء", self.names(self.search("١٠٥٠", self.viewer)))
        self.assertEqual(normalize_for_search("  إسْلامِيّة  "), "اسلاميه")

    def test_multi_word_query_matches_all_words_in_any_order(self):
        self.assertEqual(self.names(self.search("الوزراء القدس", self.viewer)), ["مجلس الوزراء"])
        self.assertEqual(self.names(self.search("مجلس غير-موجود", self.viewer)), [])

    def test_name_matches_ranked_first(self):
        Location.objects.create(place_name="شارع ارحيم", description="تل الهوا - مجلس الوزراء")
        names = self.names(self.search("مجلس", self.viewer))
        self.assertEqual(names[0], "مجلس الوزراء")
        self.assertEqual(len(names), 2)

    def test_no_results(self):
        res = self.search("لا يوجد شيء كهذا", self.viewer)
        self.assertEqual(res.data["count"], 0)

    def test_pagination(self):
        for i in range(30):
            Location.objects.create(place_name=f"موقع تجريبي {i}")
        self.client.force_authenticate(self.viewer)
        res = self.client.get(URL, {"search": "تجريبي"})
        self.assertEqual(res.data["count"], 30)
        self.assertEqual(len(res.data["results"]), 20)
        self.assertIsNotNone(res.data["next"])


class AnonymousAccessTests(LocationTestBase):
    @override_settings(PUBLIC_SEARCH_ENABLED=False)
    def test_anonymous_cannot_search_when_public_search_disabled(self):
        self.assertEqual(self.client.get(URL).status_code, status.HTTP_403_FORBIDDEN)

    @override_settings(PUBLIC_SEARCH_ENABLED=True)
    def test_anonymous_can_search_when_enabled_but_not_see_editors(self):
        res = self.client.get(URL, {"search": "مجلس"})
        self.assertEqual(res.status_code, 200)
        self.assertNotIn("updated_by_name", res.data["results"][0])

    @override_settings(PUBLIC_SEARCH_ENABLED=True)
    def test_anonymous_can_never_mutate(self):
        self.assertEqual(self.client.post(URL, {"place_name": "x"}).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.client.patch(f"{URL}{self.council.pk}/", {"place_name": "x"}).status_code, 403)
        self.assertEqual(self.client.delete(f"{URL}{self.council.pk}/").status_code, 403)
        self.council.refresh_from_db()
        self.assertTrue(self.council.is_active)


class CrudTests(LocationTestBase):
    payload = {"place_name": "مسجد الوحدة", "description": "شارع المجادلة الشاطئ",
               "building_number": "122", "street_number": "60310"}

    def test_full_lifecycle_as_editor(self):
        self.client.force_authenticate(self.editor)
        res = self.client.post(URL, self.payload)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        pk = res.data["id"]

        found = self.client.get(URL, {"search": "مسجد الوحدة"}).data["results"]
        self.assertEqual([(r["description"], r["building_number"], r["street_number"]) for r in found],
                         [("شارع المجادلة الشاطئ", "122", "60310")])

        res = self.client.patch(f"{URL}{pk}/", {"building_number": "123"})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["building_number"], "123")

        res = self.client.put(f"{URL}{pk}/", {**self.payload, "street_number": "60311"})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["street_number"], "60311")
        self.assertEqual(res.data["building_number"], "122")

        self.assertEqual(self.client.delete(f"{URL}{pk}/").status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(self.client.get(URL, {"search": "مسجد الوحدة"}).data["count"], 0)
        self.assertEqual(self.client.get(f"{URL}{pk}/").status_code, 404)

        loc = Location.objects.get(pk=pk)  # soft delete: row still exists
        self.assertFalse(loc.is_active)
        self.assertEqual(loc.deleted_by, self.editor)
        self.assertIsNotNone(loc.deleted_at)
        self.assertEqual(loc.created_by, self.editor)

        actions = list(AuditLog.objects.filter(entity_type="location", entity_id=str(pk))
                       .order_by("id").values_list("action", "user__username"))
        self.assertEqual(actions, [("CREATE", "editor1"), ("UPDATE", "editor1"), ("UPDATE", "editor1"),
                                   ("DELETE", "editor1")])
        update = AuditLog.objects.filter(entity_id=str(pk), action="UPDATE").order_by("id").first()
        self.assertEqual(update.before_data["building_number"], "122")
        self.assertEqual(update.after_data["building_number"], "123")

    def test_viewer_cannot_mutate(self):
        self.client.force_authenticate(self.viewer)
        self.assertEqual(self.client.post(URL, self.payload).status_code, 403)
        self.assertEqual(self.client.patch(f"{URL}{self.council.pk}/", {"place_name": "x"}).status_code, 403)
        self.assertEqual(self.client.put(f"{URL}{self.council.pk}/", self.payload).status_code, 403)
        self.assertEqual(self.client.delete(f"{URL}{self.council.pk}/").status_code, 403)
        self.council.refresh_from_db()
        self.assertEqual(self.council.place_name, "مجلس الوزراء")
        self.assertTrue(self.council.is_active)

    def test_inactive_editor_cannot_mutate(self):
        self.editor.is_active = False
        self.editor.save()
        self.client.force_authenticate(self.editor)
        self.assertEqual(self.client.post(URL, self.payload).status_code, 403)

    def test_duplicates_allowed(self):
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.post(URL, self.payload).status_code, 201)
        self.assertEqual(self.client.post(URL, self.payload).status_code, 201)
        self.assertEqual(self.client.get(URL, {"search": "مسجد الوحدة"}).data["count"], 2)

    def test_validation(self):
        self.client.force_authenticate(self.editor)
        res = self.client.post(URL, {"place_name": "   "})
        self.assertEqual(res.status_code, 400)
        self.assertIn("place_name", res.data)
        res = self.client.post(URL, {"place_name": "x", "building_number": "<script>"})
        self.assertEqual(res.status_code, 400)
        self.assertIn("building_number", res.data)
        res = self.client.post(URL, {"place_name": "x", "street_number": "1" * 21})
        self.assertEqual(res.status_code, 400)

    def test_whitespace_trimmed_and_leading_zeros_preserved(self):
        self.client.force_authenticate(self.editor)
        res = self.client.post(URL, {"place_name": "  مفترق   تجريبي  ", "building_number": "00123",
                                     "street_number": " 0050 "})
        self.assertEqual(res.status_code, 201)
        self.assertEqual(res.data["place_name"], "مفترق تجريبي")
        self.assertEqual(res.data["building_number"], "00123")
        self.assertEqual(res.data["street_number"], "0050")

    def test_admin_can_list_and_restore_deleted(self):
        self.client.force_authenticate(self.admin)
        self.client.delete(f"{URL}{self.council.pk}/")
        res = self.client.get(URL, {"status": "deleted"})
        self.assertEqual([r["id"] for r in res.data["results"]], [self.council.pk])
        self.assertEqual(self.client.post(f"{URL}{self.council.pk}/restore/").status_code, 200)
        self.assertEqual(self.client.get(URL, {"search": "مجلس الوزراء"}).data["count"], 1)
        self.assertTrue(AuditLog.objects.filter(action="RESTORE", entity_id=str(self.council.pk)).exists())

    def test_editor_cannot_see_deleted_or_restore(self):
        self.council.is_active = False
        self.council.save()
        self.client.force_authenticate(self.editor)
        self.assertEqual(self.client.get(URL, {"status": "deleted"}).data["count"], 4)  # falls back to active
        self.assertEqual(self.client.post(f"{URL}{self.council.pk}/restore/").status_code, 403)


class StreetNameTests(APITestCase):
    def test_search_and_permissions(self):
        StreetName.objects.create(common_name="الشفا", official_name="عزالدين القسام")
        viewer = make_user("v", User.Role.VIEWER)
        editor = make_user("e", User.Role.EDITOR)
        self.client.force_authenticate(viewer)
        res = self.client.get("/api/streets/", {"search": "القسام"})
        self.assertEqual(res.data["results"][0]["common_name"], "الشفا")
        self.assertEqual(self.client.post("/api/streets/", {"common_name": "a", "official_name": "b"}).status_code,
                         403)
        self.client.force_authenticate(editor)
        self.assertEqual(self.client.post("/api/streets/", {"common_name": "a", "official_name": "b"}).status_code,
                         201)
        self.assertEqual(self.client.post("/api/streets/", {"common_name": "a"}).status_code, 400)


class ImportCommandTests(APITestCase):
    """The migration command reads a workbook laid out like the legacy file."""

    def build_workbook(self, directory):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "فارغة"
        ws = wb.create_sheet("ورقة1")
        ws.merge_cells("E2:H3")
        ws["E2"] = "SULAIMAN  MOEN  HABIB"
        for col, text in zip("BCEFGH", ["الشارع العام ", "الشارع الرسمي", "التقاطع", "المسمي العام", "مبني", "شارع"]):
            ws[f"{col}5"] = text
        rows = [
            ("الشفا", "عزالدين القسام", "مفترق الشعبية", "عمارة لظن", 22, 1760),
            ("اللببيدي", "خليل الوزير", "مجلس الوزراء", "تل الهوا / شارع القدس", 95, 1050),
            (None, None, "مدرسة بنات الصبرة", "مدرسة الصبرة الإعدادية المشتركة", "55F", 2456),
            (None, None, " حارة ام الارانب", " الزيتون  -  مربع ", 87.0, 2207),
            (None, None, "مكرر", "نفس الوصف", 1, 2),
            (None, None, "مكرر", "نفس الوصف", 1, 2),
        ]
        for i, (b, c, e, f, g, h) in enumerate(rows, start=6):
            ws[f"B{i}"], ws[f"C{i}"], ws[f"E{i}"], ws[f"F{i}"], ws[f"G{i}"], ws[f"H{i}"] = b, c, e, f, g, h
        ws.merge_cells("B20:C25")
        ws["B20"] = "ملاحظة إرشادية طويلة\nليست سجلاً"
        ws["J6"] = "أثر جميل"
        ws["J10"] = "جسر"
        ws["J11"] = '=IF($J$10="","",IFERROR(INDEX($F:$F,MATCH("*"&$J$10&"*",$E:$E,0)),""))'
        ws["J40"], ws["K40"] = "=F6", "=H6"
        path = Path(directory) / "legacy.xlsx"
        wb.save(path)
        return path

    def test_import_is_accurate_idempotent_and_leaves_source_untouched(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self.build_workbook(tmp)
            before = path.read_bytes()
            out = StringIO()
            call_command("import_locations", str(path), stdout=out)
            self.assertIn("All source records verified", out.getvalue())
            self.assertEqual(path.read_bytes(), before)

            self.assertEqual(Location.objects.count(), 6)
            school = Location.objects.get(place_name="مدرسة بنات الصبرة")
            self.assertEqual(school.building_number, "55F")
            council = Location.objects.get(place_name="مجلس الوزراء")
            self.assertEqual((council.description, council.building_number, council.street_number),
                             ("تل الهوا / شارع القدس", "95", "1050"))
            hara = Location.objects.get(place_name="حارة ام الارانب")
            self.assertEqual((hara.description, hara.building_number), ("الزيتون - مربع", "87"))
            self.assertEqual(Location.objects.filter(place_name="مكرر").count(), 2)  # legit duplicates kept
            self.assertEqual(StreetName.objects.count(), 2)  # merged note skipped
            self.assertFalse(Location.objects.filter(place_name__startswith="=").exists())

            call_command("import_locations", str(path), stdout=StringIO())  # re-run: no new rows
            self.assertEqual(Location.objects.count(), 6)
            self.assertEqual(StreetName.objects.count(), 2)

            res_user = make_user("v", User.Role.VIEWER)
            self.client.force_authenticate(res_user)
            self.assertEqual(self.client.get(URL, {"search": "مجلس"}).data["count"], 1)

    def test_dry_run_writes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self.build_workbook(tmp)
            call_command("import_locations", str(path), "--dry-run", stdout=StringIO())
        self.assertEqual(Location.objects.count(), 0)
